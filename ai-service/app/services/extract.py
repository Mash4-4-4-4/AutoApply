import io

MAX_BYTES = 5 * 1024 * 1024


class ExtractionError(Exception):
    pass


def detect_kind(data: bytes, filename: str) -> str:
    name = (filename or "").lower()
    if data[:4] == b"%PDF":
        return "pdf"
    if data[:2] == b"PK" and name.endswith(".docx"):
        return "docx"
    raise ExtractionError("Unsupported file. Upload a PDF or DOCX resume.")


def extract_text(data: bytes, kind: str) -> str:
    try:
        text = _pdf(data) if kind == "pdf" else _docx(data)
    except ExtractionError:
        raise
    except Exception as exc:  # corrupt / encrypted files
        raise ExtractionError("Could not open this file. It may be corrupted or password-protected.") from exc
    text = text.strip()
    if len(text) < 30:
        raise ExtractionError("No readable text found. Scanned/image-only resumes are not supported yet.")
    return text


def _pdf(data: bytes) -> str:
    import pymupdf as fitz

    with fitz.open(stream=data, filetype="pdf") as doc:
        if doc.needs_pass:
            raise ExtractionError("This PDF is password-protected.")
        return "\n".join(page.get_text() for page in list(doc)[:10])


def _docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(lines)
